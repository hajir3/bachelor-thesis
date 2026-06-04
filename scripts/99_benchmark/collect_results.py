#!/usr/bin/env python3
"""collect_results.py — Append benchmark timings to an XLSX file.

Usage (called by benchmark.sh):
    python3 collect_results.py \
        --output benchmark_results.xlsx \
        --codename bookworm \
        --total 12.345 \
        --timings "download_packages:1.23,parse_packages:2.34,..."
"""
import argparse
import os
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

# All possible step names in pipeline order
ALL_STEPS = [
    "download_packages",
    "download_sources_idx",
    "parse_packages",
    "parse_sources",
    "find_persistent",
    "download_tarballs",
    "analyze_installed_size",
    "analyze_dependencies",
    "analyze_tokei",
    "aggregate_tokei",
    "analyze_patches",
    "analyze_complexity",
]

HEADER = ["timestamp", "codename"] + ALL_STEPS + ["total"]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, help="Path to XLSX file")
    parser.add_argument("--codename", required=True)
    parser.add_argument("--total", required=True)
    parser.add_argument("--timings", required=True,
                        help="Comma-separated step:seconds pairs")
    args = parser.parse_args()

    # Parse timings into a dict
    timings = {}
    for pair in args.timings.split(","):
        name, secs = pair.split(":")
        timings[name] = float(secs)

    # Open or create workbook
    if os.path.exists(args.output):
        wb = load_workbook(args.output)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Benchmark Results"
        ws.append(HEADER)
        for cell in ws[1]:
            cell.font = Font(bold=True)

    # Build row
    row = [
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        args.codename,
    ]
    for step in ALL_STEPS:
        row.append(timings.get(step, None))
    row.append(float(args.total))

    ws.append(row)
    wb.save(args.output)


if __name__ == "__main__":
    main()
